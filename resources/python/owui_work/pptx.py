"""PowerPoint presentations (python-pptx).

    from owui_work.pptx import build_presentation

    build_presentation("deck.pptx", [
        {"layout": "title", "title": "Q3 results", "subtitle": "Sales team"},
        {"title": "Highlights", "bullets": ["Revenue +12%", {"text": "North +20%", "level": 1}],
         "notes": "Mention the new contract"},
        {"title": "By region", "table": {"columns": ["Region", "Q3"], "rows": [["North", 140]]}},
        {"title": "Trend", "chart": {"type": "line", "categories": ["Jul", "Aug", "Sep"],
                                      "series": [{"name": "Sales", "values": [10, 12, 15]}]}},
        {"title": "Compare", "left": ["Plan"], "right": ["Actual"]},
        {"title": "Photo", "image": "chart.png"},
    ])
"""

import os

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.util import Emu, Inches, Pt

_CHARTS = {
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "pie": XL_CHART_TYPE.PIE,
}

# Default template layouts
_TITLE, _CONTENT, _TWO, _TITLE_ONLY = 0, 1, 3, 5


def _layout(prs, index):
    layouts = prs.slide_layouts
    return layouts[index] if index < len(layouts) else layouts[-1]


def _fill_bullets(frame, bullets, size=None):
    frame.clear()
    first = True
    for item in bullets:
        text, level = (item.get("text", ""), int(item.get("level", 0))) if isinstance(item, dict) else (str(item), 0)
        p = frame.paragraphs[0] if first else frame.add_paragraph()
        p.text = text
        p.level = max(0, min(level, 4))
        if size:
            for run in p.runs:
                run.font.size = Pt(size)
        first = False


def _body_box(prs, slide):
    """Area below the title for tables, charts and pictures."""
    top = Inches(1.5)
    title = slide.shapes.title
    if title is not None:
        top = max(top, title.top + title.height + Inches(0.2))
    left = Inches(0.6)
    return left, top, prs.slide_width - 2 * left, prs.slide_height - top - Inches(0.5)


def build_presentation(path, slides, template=None, widescreen=True):
    """Create a .pptx from a list of slide dicts and return its path.

    Slide keys: title, subtitle (with layout "title"), bullets, left/right
    (two columns), table {columns, rows}, chart {type: bar|column|line|pie,
    categories, series: [{name, values}]}, image (file path), notes.
    template: an existing .pptx whose theme and layouts are used.
    """
    prs = Presentation(template) if template else Presentation()
    if not template and widescreen:
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

    for spec in slides:
        if spec.get("layout") == "title":
            slide = prs.slides.add_slide(_layout(prs, _TITLE))
            slide.shapes.title.text = spec.get("title", "")
            if len(slide.placeholders) > 1:
                slide.placeholders[1].text = spec.get("subtitle", "")
        elif spec.get("left") is not None or spec.get("right") is not None:
            slide = prs.slides.add_slide(_layout(prs, _TWO))
            slide.shapes.title.text = spec.get("title", "")
            bodies = [p for p in slide.placeholders if p.placeholder_format.idx != 0]
            if len(bodies) >= 2:
                _fill_bullets(bodies[0].text_frame, spec.get("left") or [])
                _fill_bullets(bodies[1].text_frame, spec.get("right") or [])
        elif spec.get("bullets") is not None:
            slide = prs.slides.add_slide(_layout(prs, _CONTENT))
            slide.shapes.title.text = spec.get("title", "")
            body = [p for p in slide.placeholders if p.placeholder_format.idx != 0]
            if body:
                _fill_bullets(body[0].text_frame, spec["bullets"], spec.get("font_size"))
        else:
            slide = prs.slides.add_slide(_layout(prs, _TITLE_ONLY))
            if slide.shapes.title is not None:
                slide.shapes.title.text = spec.get("title", "")
            left, top, width, height = _body_box(prs, slide)

            if spec.get("table"):
                table_spec = spec["table"]
                rows = [table_spec.get("columns", [])] + list(table_spec.get("rows", []))
                cols = max(len(r) for r in rows)
                shape = slide.shapes.add_table(len(rows), cols, left, top, width, Emu(min(height, Inches(0.4) * len(rows))))
                for r, row in enumerate(rows):
                    for c in range(cols):
                        shape.table.cell(r, c).text = str(row[c]) if c < len(row) and row[c] is not None else ""
            elif spec.get("chart"):
                chart_spec = spec["chart"]
                data = CategoryChartData()
                data.categories = chart_spec.get("categories", [])
                for series in chart_spec.get("series", []):
                    data.add_series(series.get("name", ""), series.get("values", []))
                kind = _CHARTS.get(chart_spec.get("type", "column"), XL_CHART_TYPE.COLUMN_CLUSTERED)
                chart = slide.shapes.add_chart(kind, left, top, width, height, data).chart
                chart.has_legend = len(chart_spec.get("series", [])) > 1 or kind == XL_CHART_TYPE.PIE
                if chart.has_legend:
                    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
                    chart.legend.include_in_layout = False
                if chart_spec.get("title"):
                    chart.has_title = True
                    chart.chart_title.text_frame.text = chart_spec["title"]
            elif spec.get("image"):
                image = spec["image"]
                if os.path.exists(image):
                    picture = slide.shapes.add_picture(image, left, top)
                    # Fit inside the body area, keep aspect ratio
                    scale = min(width / picture.width, height / picture.height, 1)
                    picture.width, picture.height = int(picture.width * scale), int(picture.height * scale)
                    picture.left = int(left + (width - picture.width) / 2)
                else:
                    box = slide.shapes.add_textbox(left, top, width, Inches(1))
                    box.text_frame.text = f"[image not found: {image}]"
            elif spec.get("text"):
                box = slide.shapes.add_textbox(left, top, width, height)
                box.text_frame.word_wrap = True
                box.text_frame.text = spec["text"]

        if spec.get("notes"):
            slide.notes_slide.notes_text_frame.text = spec["notes"]

    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    prs.save(path)
    return path
