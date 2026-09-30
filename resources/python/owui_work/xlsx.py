"""Excel workbooks (openpyxl).

    from owui_work.xlsx import write_workbook, update_cells, append_rows

    write_workbook("sales.xlsx", [
        {
            "name": "Sales",
            "rows": [
                ["Region", "Q1", "Q2", "Total"],
                ["North", 120, 140, "=B2+C2"],
                ["South", 90, 110, "=B3+C3"],
                ["Total", "=SUM(B2:B3)", "=SUM(C2:C3)", "=SUM(D2:D3)"],
            ],
            "number_formats": {"B:D": "#,##0"},
            "chart": {"type": "bar", "title": "Sales by region",
                      "data": "B1:C3", "categories": "A2:A3", "anchor": "F2"},
        }
    ])

Strings starting with "=" are written as formulas; Excel calculates them
when the file is opened.
"""

import os
import re

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import column_index_from_string, get_column_letter, range_boundaries

_HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")


def _col_span(spec):
    """'B' or 'B:D' → list of column indexes."""
    if ":" in spec:
        a, b = spec.split(":", 1)
        return list(range(column_index_from_string(a), column_index_from_string(b) + 1))
    return [column_index_from_string(spec)]


def _auto_width(ws, max_width=60):
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            text = str(cell.value)
            if text.startswith("="):
                text = "0000000000"
            longest = max(len(part) for part in text.split("\n"))
            widths[cell.column] = max(widths.get(cell.column, 0), longest)
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = min(max(width + 2, 8), max_width)


def _add_chart(ws, spec):
    kind = spec.get("type", "bar")
    chart = {"bar": BarChart, "column": BarChart, "line": LineChart, "pie": PieChart}.get(kind, BarChart)()
    if kind == "bar":
        chart.type = "bar"
    elif kind == "column":
        chart.type = "col"
    if spec.get("title"):
        chart.title = spec["title"]
    min_col, min_row, max_col, max_row = range_boundaries(spec["data"])
    data = Reference(ws, min_col=min_col, min_row=min_row, max_col=max_col, max_row=max_row)
    chart.add_data(data, titles_from_data=spec.get("titles_from_data", True))
    if spec.get("categories"):
        c1, r1, c2, r2 = range_boundaries(spec["categories"])
        chart.set_categories(Reference(ws, min_col=c1, min_row=r1, max_col=c2, max_row=r2))
    chart.width = spec.get("width", 16)
    chart.height = spec.get("height", 8)
    ws.add_chart(chart, spec.get("anchor", "H2"))


def _fill_sheet(ws, sheet):
    rows = sheet.get("rows")
    if rows is None and "columns" in sheet:
        rows = [sheet["columns"]] + [list(r) for r in sheet.get("data", [])]
    rows = rows or []
    for row in rows:
        ws.append(list(row))

    header = sheet.get("header", True) and bool(rows)
    if header:
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = _HEADER_FILL
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.freeze_panes = sheet.get("freeze", "A2")
        if sheet.get("autofilter", True) and ws.max_row > 1:
            ws.auto_filter.ref = ws.dimensions
    elif sheet.get("freeze"):
        ws.freeze_panes = sheet["freeze"]

    for spec, fmt in (sheet.get("number_formats") or {}).items():
        cols = _col_span(spec)
        for row in ws.iter_rows(min_row=2 if header else 1):
            for cell in row:
                if cell.column in cols:
                    cell.number_format = fmt

    widths = sheet.get("widths")
    if widths:
        for col, width in widths.items():
            ws.column_dimensions[col].width = width
    else:
        _auto_width(ws)

    charts = sheet.get("charts") or ([sheet["chart"]] if sheet.get("chart") else [])
    for spec in charts:
        _add_chart(ws, spec)


def write_workbook(path, sheets):
    """Create a new .xlsx from a list of sheet dicts and return its path.

    Sheet keys: name, rows (list of lists — the first row is the header) or
    columns + data, header (default True: bold, frozen, filter),
    number_formats ({"B:D": "#,##0.00"}), widths ({"A": 30}), freeze,
    chart / charts ({type: bar|column|line|pie, title, data, categories,
    anchor}).
    """
    wb = Workbook()
    wb.remove(wb.active)
    for index, sheet in enumerate(sheets):
        name = re.sub(r"[\[\]:*?/\\]", "_", str(sheet.get("name") or f"Sheet{index + 1}"))[:31]
        _fill_sheet(wb.create_sheet(name), sheet)
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    wb.save(path)
    return path


def update_cells(path, cells, sheet=None, save_as=None):
    """Set cells in an existing workbook: update_cells("a.xlsx", {"B2": 10, "C5": "=SUM(C1:C4)"}).

    Note: openpyxl drops charts and images of the existing file when it
    saves; use save_as to write a copy if the workbook has them.
    """
    wb = load_workbook(path)
    ws = wb[sheet] if sheet else wb.active
    for ref, value in cells.items():
        ws[ref] = value
    wb.save(save_as or path)
    return save_as or path


def append_rows(path, rows, sheet=None, save_as=None):
    """Append rows at the end of a sheet in an existing workbook."""
    wb = load_workbook(path)
    ws = wb[sheet] if sheet else wb.active
    for row in rows:
        ws.append(list(row))
    wb.save(save_as or path)
    return save_as or path
